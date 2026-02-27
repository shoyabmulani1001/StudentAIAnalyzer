"""
excel_reader.py
Reads all Excel subject files from the excel_files/ directory.
Extracts student internal (CO Int. Attn) and external (CO Ext. Attn) marks,
merges them, and returns structured student-subject records.
"""

import os
import re
import openpyxl
from typing import List, Dict, Any


EXCEL_DIR = os.path.join(os.path.dirname(__file__), "excel_files")

# Sheet names
INT_SHEET  = "CO Int. Attn"
EXT_SHEET  = "CO Ext. Attn"
DETAILS_SHEET = "SubjectDetails"

# Max marks constants removed since we detect dynamically


# ── Subject name ──────────────────────────────────────────────────────────────

def _get_subject_name(wb: openpyxl.Workbook, filename: str) -> str:
    """Extract subject name from SubjectDetails sheet, fallback to filename."""
    try:
        ws = wb[DETAILS_SHEET]
        for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
            if row[0] and str(row[0]).strip().lower().startswith("subject"):
                val = row[1] if len(row) > 1 else None
                if val:
                    return str(val).strip()
    except Exception:
        pass
    name = os.path.splitext(filename)[0]
    parts = name.split(" ", 1)
    return parts[1] if len(parts) > 1 else name


# ── Internal sheet ────────────────────────────────────────────────────────────

def _find_total_col(ws) -> tuple:
    """
    Dynamically find the column index of 'Total Marks Obtained' and the
    raw maximum (e.g., 25, 40, etc.).
    """
    for row in ws.iter_rows(max_row=15, values_only=True):
        for idx, cell in enumerate(row):
            if cell and isinstance(cell, str) and "total marks obtained" in cell.lower():
                raw_max = 25 # default if not found
                m = re.search(r"\((\d+)\)", cell)
                if m:
                    raw_max = int(m.group(1))
                return idx, raw_max
            elif cell and isinstance(cell, str) and ("out of" in cell.lower() or "max marks" in cell.lower()):
                raw_max = 25
                m = re.search(r"(\d+)", cell)
                if m:
                    raw_max = int(m.group(1))
                return idx, raw_max
    return None, 25


def _read_internal(ws) -> tuple:
    """
    Parse CO Int. Attn sheet.
    Dynamically detects the 'Total Marks Obtained' column and maximum marks.
    Returns tuple: ({ STUDENT_NAME -> {seat_no, internal_marks} }, max_marks)
    """
    total_col, raw_max = _find_total_col(ws)
    students = {}
    data_started = False

    for row in ws.iter_rows(values_only=True):
        if not data_started:
            if row[0] in ("Sr. No", "Sr.No"):
                data_started = True
                continue
            elif isinstance(row[0], int) and row[0] == 1:
                data_started = True   # fall-through to process this row
            else:
                continue

        if row[0] is None:
            continue
        try:
            int(row[0])
        except (ValueError, TypeError):
            continue

        name    = str(row[2]).strip().upper() if row[2] else None
        seat_no = str(row[1]).strip()         if row[1] else None

        # Try dynamically detected column first, then common fallback indices
        internal = None
        search_cols = ([total_col] if total_col is not None else []) + [20, 23, 19, 21, 22, 24]
        for col_idx in search_cols:
            if col_idx is not None and col_idx < len(row) and row[col_idx] is not None:
                try:
                    val = float(row[col_idx])
                    if val > raw_max:
                        continue      # out of expected range — skip
                    internal = val
                    break
                except (ValueError, TypeError):
                    pass

        if name and internal is not None:
            students[name] = {"seat_no": seat_no, "internal_marks": round(internal, 2)}

    return students, raw_max


# ── External sheet ────────────────────────────────────────────────────────────

def _read_external(ws) -> tuple:
    """
    Parse CO Ext. Attn sheet.
    Columns: 0=Roll, 1=SeatNo, 2=Name, 3=Marks
    Detects max marks from headers, default 50.
    Returns tuple: ({ STUDENT_NAME -> {seat_no, external_marks} }, max_marks)
    """
    max_marks = 50
    # Search for max marks in the first few rows
    for row in ws.iter_rows(max_row=5, values_only=True):
         for cell in row:
             if cell and isinstance(cell, str) and ("out of" in cell.lower() or "max" in cell.lower() or "total" in cell.lower()):
                 m = re.search(r"\((\d+)\)", cell)
                 if not m:
                     m = re.search(r"(\d+)", cell)
                 if m:
                     max_marks = int(m.group(1))
                     break

    students = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] is None:
            continue
        try:
            int(row[0])
        except (ValueError, TypeError):
            continue

        name    = str(row[2]).strip().upper() if row[2] else None
        seat_no = str(row[1]).strip()         if row[1] else None

        external = None
        if len(row) > 3 and row[3] is not None:
            try:
                external = float(row[3])
                if external > max_marks:
                     # Attempt to find it in other columns if index 3 is out of bounds
                     for col_idx in range(4, len(row)):
                         if row[col_idx] is not None:
                             val = float(row[col_idx])
                             if val <= max_marks:
                                 external = val
                                 break
            except (ValueError, TypeError):
                pass

        if name and external is not None and external <= max_marks:
            students[name] = {"seat_no": seat_no, "external_marks": round(external, 2)}

    return students, max_marks


# ── Public API ────────────────────────────────────────────────────────────────

def read_all_files() -> List[Dict[str, Any]]:
    """
    Read all .xlsx files from excel_files/ and return a flat list of records:
    [
      {
        name, seat_no, subject,
        obtained_marks, max_marks,
        percentage, file
      }, ...
    ]
    """
    records = []
    if not os.path.isdir(EXCEL_DIR):
        print(f"[ERROR] excel_files directory not found: {EXCEL_DIR}")
        return records

    files = sorted(f for f in os.listdir(EXCEL_DIR) if f.endswith(".xlsx"))
    
    # Map normalized canonical name parts to the first seen display name
    global_name_map = {}

    for filename in files:
        filepath = os.path.join(EXCEL_DIR, filename)
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
        except Exception as e:
            print(f"[WARN] Could not open {filename}: {e}")
            continue

        subject = _get_subject_name(wb, filename)

        if INT_SHEET not in wb.sheetnames or EXT_SHEET not in wb.sheetnames:
            print(f"[WARN] Missing sheets in {filename}: {wb.sheetnames}")
            continue

        internal_data, internal_max = _read_internal(wb[INT_SHEET])
        external_data, external_max = _read_external(wb[EXT_SHEET])

        # If data exists for both internal and external, add them. Otherwise just the available max
        has_internal = len(internal_data) > 0
        has_external = len(external_data) > 0

        # Calculate subject max dynamically
        if has_internal and has_external:
             subject_max = internal_max + external_max
        elif has_internal:
             subject_max = internal_max
        elif has_external:
             subject_max = external_max
        else:
             subject_max = internal_max + external_max

        print(f"{subject} -> {subject_max}")

        all_names = set(internal_data.keys()) | set(external_data.keys())

        for name in all_names:
            int_info = internal_data.get(name, {})
            ext_info = external_data.get(name, {})

            internal = int_info.get("internal_marks")
            external = ext_info.get("external_marks")
            seat_no  = int_info.get("seat_no") or ext_info.get("seat_no")

            if internal is None and external is None:
                continue

            internal = internal if internal is not None else 0.0
            external = external if external is not None else 0.0

            total = internal + external
            percentage = round((total / subject_max) * 100, 2) if subject_max > 0 else 0
            
            # Create a canonical key to resolve names like "AJAYKUMAR PHAD" vs "PHAD AJAYKUMAR"
            canonical_key = tuple(sorted(name.split()))
            if canonical_key not in global_name_map:
                # Prefer names that likely start with Surname (often the majority format in Indian universities)
                # But to keep it simple, just use the first format encountered as the consistent name
                global_name_map[canonical_key] = name
            
            display_name = global_name_map[canonical_key]

            records.append({
                "name":           display_name,
                "seat_no":        seat_no,
                "subject":        subject,
                "obtained_marks": total,
                "max_marks":      subject_max,
                "percentage":     percentage,
                "file":           filename,
            })

        wb.close()

    return records


def get_all_students(records: List[Dict]) -> List[str]:
    return sorted(set(r["name"] for r in records))


def get_all_subjects(records: List[Dict]) -> List[str]:
    return sorted(set(r["subject"] for r in records))
