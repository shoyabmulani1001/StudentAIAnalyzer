import os
import sys
import openpyxl

EXCEL_DIR = "backend/excel_files"
files = [f for f in os.listdir(EXCEL_DIR) if f.endswith('.xlsx')]

for f in files:
    if "Object Oriented" in f or "OOSE" in f.upper():
        print(f"Reading {f}")
        wb = openpyxl.load_workbook(os.path.join(EXCEL_DIR, f), data_only=True)
        if "CO Int. Attn" in wb.sheetnames:
            ws = wb["CO Int. Attn"]
            count = 0
            for row in ws.iter_rows(values_only=True):
                # Just print the first 3 columns of rows where row[0] is an int
                if isinstance(row[0], int):
                    print(row[:3])
                    count += 1
            print(f"Found {count} integer rows in internal.")
        
        if "CO Ext. Attn" in wb.sheetnames:
            ws = wb["CO Ext. Attn"]
            count = 0
            for row in ws.iter_rows(values_only=True):
                if isinstance(row[0], int):
                    print(row[:3])
                    count += 1
            print(f"Found {count} integer rows in external.")
        
        wb.close()
