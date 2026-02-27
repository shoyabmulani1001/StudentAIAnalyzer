import os
import openpyxl

EXCEL_DIR = "backend/excel_files"
files = [f for f in os.listdir(EXCEL_DIR) if f.endswith('.xlsx')]

for f in files:
    wb = openpyxl.load_workbook(os.path.join(EXCEL_DIR, f), data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str) and "PHAD" in cell.upper() and "AJAY" in cell.upper():
                    print(f"Found 'Phad Ajaykumar' in {f} ({sheet})")
    wb.close()
